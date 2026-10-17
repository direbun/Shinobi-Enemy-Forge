"""Extract static Naruto 5e adversary data from the supplied workbook.

The workbook is treated as untrusted input. This script reads cached values only,
does not execute formulas, and never follows workbook links.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.replace("\u00a0", " ").strip()
    return value


def slug(value: str) -> str:
    return "-".join(
        "".join(character.lower() if character.isalnum() else " " for character in value).split()
    )


def unique(values):
    result = []
    seen = set()
    for value in values:
        value = clean(value)
        key = str(value).casefold()
        if value and key not in seen:
            seen.add(key)
            result.append(value)
    return result


def rows_with_name(sheet, start, end, name_column, columns):
    rows = []
    for row in range(start, end + 1):
        name = clean(sheet.cell(row, name_column).value)
        if not name:
            continue
        item = {key: clean(sheet.cell(row, column).value) for key, column in columns.items()}
        item["id"] = f"{slug(str(name))}-{row}"
        rows.append(item)
    return rows


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: extract_enemy_rules.py INPUT.xlsx OUTPUT.json")

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    workbook = load_workbook(input_path, data_only=True, read_only=False, keep_links=False)

    charts = workbook["ChartsMain"]
    levels = []
    for row in range(3, 37):
        level = clean(charts.cell(row, 2).value)
        if level == "":
            continue
        levels.append(
            {
                "level": level,
                "baseAc": clean(charts.cell(row, 3).value),
                "proficiency": clean(charts.cell(row, 4).value),
                "attackBonus": clean(charts.cell(row, 7).value),
                "damagePerAction": clean(charts.cell(row, 8).value),
                "standardArray": [clean(charts.cell(row, col).value) for col in range(9, 15)],
                "eliteArray": [clean(charts.cell(row, col).value) for col in range(16, 22)],
                "traitLimits": {
                    "Minion": clean(charts.cell(row, 24).value),
                    "Standard": clean(charts.cell(row, 25).value),
                    "Elite": clean(charts.cell(row, 26).value),
                    "Solo": clean(charts.cell(row, 27).value),
                },
            }
        )

    roles = []
    for row in range(3, 14):
        name = clean(charts.cell(row, 28).value)
        if not name or name == "(Base)":
            continue
        roles.append(
            {
                "id": slug(name),
                "name": name,
                "ac": clean(charts.cell(row, 29).value),
                "save": clean(charts.cell(row, 30).value),
                "hp": clean(charts.cell(row, 31).value),
                "chakra": clean(charts.cell(row, 32).value),
                "damage": clean(charts.cell(row, 33).value),
                "speed": clean(charts.cell(row, 34).value),
            }
        )

    ranks = []
    for row in range(3, 9):
        name = clean(charts.cell(row, 43).value)
        if not name:
            continue
        ranks.append(
            {
                "name": name,
                "hitDie": clean(charts.cell(row, 44).value),
                "hitDieMax": clean(charts.cell(row, 45).value),
                "hitDieAverage": clean(charts.cell(row, 46).value),
                "chakraDie": clean(charts.cell(row, 47).value),
                "chakraDieMax": clean(charts.cell(row, 48).value),
                "chakraDieAverage": clean(charts.cell(row, 49).value),
                "defensiveBonus": clean(charts.cell(row, 50).value),
                "toHitBonus": clean(charts.cell(row, 51).value),
                "saveDcBonus": clean(charts.cell(row, 52).value),
                "bonusDamage": clean(charts.cell(row, 53).value),
                "bonusAbilityPoints": clean(charts.cell(row, 54).value),
            }
        )

    traits_sheet = workbook["TraitsMain"]
    traits = rows_with_name(
        traits_sheet,
        3,
        1258,
        2,
        {
            "source": 1,
            "name": 2,
            "rank": 3,
            "description": 4,
            "categories": 5,
            "type": 6,
            "rankNumber": 9,
        },
    )
    traits = [trait for trait in traits if trait["description"]]

    jutsu_sheet = workbook["JutsuMain"]
    jutsu = rows_with_name(
        jutsu_sheet,
        1,
        1820,
        1,
        {
            "name": 1,
            "discipline": 2,
            "rank": 3,
            "castingTime": 4,
            "range": 5,
            "duration": 6,
            "components": 7,
            "cost": 8,
            "keywords": 9,
            "description": 10,
            "higherRanks": 11,
        },
    )

    weapons_sheet = workbook["WeapMain"]
    weapons = rows_with_name(
        weapons_sheet,
        5,
        1000,
        1,
        {
            "name": 1,
            "damage": 2,
            "damageType": 3,
            "qualities": 4,
            "group": 5,
            "ability": 6,
            "attackType": 7,
            "range": 8,
            "targets": 9,
        },
    )

    list_columns = {
        "clans": 10,
        "roleNames": 11,
        "affiliations": 12,
        "classMods": 13,
    }
    lists = {
        key: unique(traits_sheet.cell(row, column).value for row in range(3, 1259))
        for key, column in list_columns.items()
    }

    skills = []
    for row in range(4, 25):
        name = clean(weapons_sheet.cell(row, 14).value)
        if name:
            skills.append({"name": name, "ability": clean(weapons_sheet.cell(row, 15).value)})

    data = {
        "source": {
            "title": "Copy of N5E 2025 Adversary Builder (Make a copy)",
            "version": "1.0.7",
            "workbook": input_path.name,
            "guide": "Official Naruto 5e Kage Guide.pdf",
        },
        "levels": levels,
        "roles": roles,
        "ranks": ranks,
        "classes": ["Minion", "Standard", "Elite", "Solo"],
        "lists": lists,
        "skills": skills,
        "traits": traits,
        "jutsu": jutsu,
        "weapons": weapons,
        "encounterRules": {
            "classWeights": {"Minion": 0.5, "Standard": 1, "Elite": 2, "Solo": 4},
            "classModLevelBonuses": {"None": 0, "Tier 1": 2, "Tier 2": 4, "Tier 3": 6},
            "difficulty": [
                {"maxDelta": -3, "name": "Very Easy"},
                {"maxDelta": -2, "name": "Easy"},
                {"maxDelta": 1, "name": "Moderate"},
                {"maxDelta": 2, "name": "Hard"},
                {"maxDelta": 999, "name": "Deadly"},
            ],
        },
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(
        f"Wrote {output_path}: {len(traits)} traits, {len(jutsu)} jutsu, "
        f"{len(weapons)} weapons, {len(levels)} level rows"
    )


if __name__ == "__main__":
    main()
